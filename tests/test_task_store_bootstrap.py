import datetime as dt
from pathlib import Path

import pytest

pytest.importorskip("openpyxl")

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from backend.backend import META_ID_COLUMN, PARENT_ID_COLUMN, TASK_COLUMNS, TaskStore


def _build_initial_workbook(path: Path, sheet_name: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.delete_rows(1, ws.max_row)
    ws.append(TASK_COLUMNS)
    ws.append([
        "未着手",
        "タスクA",
        "Alice",
        "高",
        dt.date(2024, 1, 1),
        "最初のタスク",
    ])
    ws.append([
        "進行中",
        "タスクB",
        "Bob",
        "中",
        dt.date(2024, 1, 2),
        "二番目のタスク",
    ])
    wb.save(path)


def test_bootstrap_persists_ids_and_respects_local_changes(tmp_path):
    excel_path = tmp_path / "board.xlsx"
    sheet_name = "Kanban"
    _build_initial_workbook(excel_path, sheet_name)

    store = TaskStore(excel_path, sheet_name=sheet_name)

    wb = load_workbook(excel_path)
    ws = wb[sheet_name]
    headers = [cell.value for cell in ws[1]]
    assert META_ID_COLUMN in headers
    id_col_idx = headers.index(META_ID_COLUMN) + 1
    assert ws.column_dimensions[get_column_letter(id_col_idx)].hidden is True
    first_id_cell = ws.cell(row=2, column=id_col_idx).value
    second_id_cell = ws.cell(row=3, column=id_col_idx).value
    assert first_id_cell
    assert second_id_cell

    df = store._df.copy()
    first_row_id = df.iloc[0][META_ID_COLUMN]
    second_row_id = df.iloc[1][META_ID_COLUMN]

    store.update_task(1, {"タスク": "ローカル修正"})
    store.delete_task(2)

    wb = load_workbook(excel_path)
    ws = wb[sheet_name]
    task_col_idx = headers.index("タスク") + 1
    ws.cell(row=2, column=task_col_idx).value = "外部編集"
    ws.cell(row=3, column=task_col_idx).value = "復活してほしくない"
    wb.save(excel_path)

    store.load_excel()

    assert len(store._df) == 1
    assert store._df.iloc[0]["タスク"] == "ローカル修正"
    assert first_row_id in store._dirty_row_ids
    assert second_row_id in store._deleted_row_ids
    assert first_row_id == store._df.iloc[0][META_ID_COLUMN]


def test_parent_id_column_exists(tmp_path):
    excel_path = tmp_path / "board.xlsx"
    _build_initial_workbook(excel_path, "Sheet1")
    store = TaskStore(excel_path)
    assert PARENT_ID_COLUMN in store._df.columns


def test_add_task_with_parent(tmp_path):
    excel_path = tmp_path / "board.xlsx"
    _build_initial_workbook(excel_path, "Sheet1")
    store = TaskStore(excel_path)
    tasks = store.get_tasks()
    parent_id = tasks[0]["__kanban_id"]
    child = store.add_task({"タスク": "子タスク", "__parent_id": parent_id})
    assert child["__parent_id"] == parent_id


def test_circular_reference_prevention(tmp_path):
    excel_path = tmp_path / "board.xlsx"
    _build_initial_workbook(excel_path, "Sheet1")
    store = TaskStore(excel_path)
    tasks = store.get_tasks()
    parent_id = tasks[0]["__kanban_id"]
    child = store.add_task({"タスク": "子タスク", "__parent_id": parent_id})
    child_id = child["__kanban_id"]
    with pytest.raises(ValueError, match="循環参照"):
        store.update_task(1, {"__parent_id": child_id})


def test_cascade_delete(tmp_path):
    excel_path = tmp_path / "board.xlsx"
    _build_initial_workbook(excel_path, "Sheet1")
    store = TaskStore(excel_path)
    tasks = store.get_tasks()
    parent_id = tasks[0]["__kanban_id"]
    store.add_task({"タスク": "子タスク1", "__parent_id": parent_id})
    store.add_task({"タスク": "子タスク2", "__parent_id": parent_id})
    assert len(store.get_tasks()) == 4
    store.delete_task(1, cascade=True)
    assert len(store.get_tasks()) == 1


def test_delete_without_cascade_promotes_children(tmp_path):
    excel_path = tmp_path / "board.xlsx"
    _build_initial_workbook(excel_path, "Sheet1")
    store = TaskStore(excel_path)
    tasks = store.get_tasks()
    parent_id = tasks[0]["__kanban_id"]
    store.add_task({"タスク": "子タスク", "__parent_id": parent_id})
    assert len(store.get_tasks()) == 3
    store.delete_task(1, cascade=False)
    remaining = store.get_tasks()
    assert len(remaining) == 2
    child = next(t for t in remaining if t["タスク"] == "子タスク")
    assert child["__parent_id"] == ""


def test_get_task_tree(tmp_path):
    excel_path = tmp_path / "board.xlsx"
    _build_initial_workbook(excel_path, "Sheet1")
    store = TaskStore(excel_path)
    tasks = store.get_tasks()
    parent_id = tasks[0]["__kanban_id"]
    store.add_task({"タスク": "子タスク", "__parent_id": parent_id})
    tree = store.get_task_tree()
    assert len(tree) >= 1
    root = next((t for t in tree if t["__kanban_id"] == parent_id), None)
    assert root is not None
    assert len(root["children"]) == 1
    assert root["children"][0]["タスク"] == "子タスク"


def test_migration_from_categories(tmp_path):
    """大分類・中分類列を持つ旧形式Excelが自動移行されることを確認"""
    excel_path = tmp_path / "old_format.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ステータス", "大分類", "中分類", "タスク", "担当者", "優先度", "期限", "備考"])
    ws.append(["未着手", "プロジェクトA", "設計", "タスク1", "Alice", "高", dt.date(2024, 1, 1), ""])
    ws.append(["進行中", "プロジェクトA", "実装", "タスク2", "Bob", "中", dt.date(2024, 1, 2), ""])
    ws.append(["未着手", "プロジェクトB", "企画", "タスク3", "Carol", "低", dt.date(2024, 1, 3), ""])
    wb.save(excel_path)

    store = TaskStore(excel_path)
    tasks = store.get_tasks()

    # 大分類・中分類列が削除されていること
    assert "大分類" not in store._df.columns
    assert "中分類" not in store._df.columns
    # __parent_id列が存在すること
    assert PARENT_ID_COLUMN in store._df.columns

    # 親タスクが自動作成されていること（プロジェクトA, プロジェクトB）
    parent_tasks = [t for t in tasks if not t["__parent_id"]]
    parent_names = [t["タスク"] for t in parent_tasks]
    assert "プロジェクトA" in parent_names
    assert "プロジェクトB" in parent_names

    # 子タスクに__parent_idが設定されていること
    child_tasks = [t for t in tasks if t["__parent_id"]]
    assert len(child_tasks) == 3

    # 中分類が備考に退避されていること
    task1 = next(t for t in tasks if t["タスク"] == "タスク1")
    assert "[旧中分類: 設計]" in task1["備考"]


def test_save_and_load_with_parent_id(tmp_path):
    excel_path = tmp_path / "board.xlsx"
    _build_initial_workbook(excel_path, "Sheet1")
    store = TaskStore(excel_path)
    tasks = store.get_tasks()
    parent_id = tasks[0]["__kanban_id"]
    store.add_task({"タスク": "子タスク", "__parent_id": parent_id})
    store.save_excel()

    store2 = TaskStore(excel_path)
    tasks2 = store2.get_tasks()
    assert len(tasks2) == 3
    child = next(t for t in tasks2 if t["タスク"] == "子タスク")
    assert child["__parent_id"] == parent_id


def test_parent_id_hidden_in_excel(tmp_path):
    excel_path = tmp_path / "board.xlsx"
    _build_initial_workbook(excel_path, "Sheet1")
    store = TaskStore(excel_path)
    store.save_excel()

    wb = load_workbook(excel_path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    assert PARENT_ID_COLUMN in headers
    parent_col_idx = headers.index(PARENT_ID_COLUMN) + 1
    assert ws.column_dimensions[get_column_letter(parent_col_idx)].hidden is True


def test_auto_sync_parent_status_when_children_complete(tmp_path):
    excel_path = tmp_path / "board.xlsx"
    _build_initial_workbook(excel_path, "Sheet1")
    store = TaskStore(excel_path)

    tasks = store.get_tasks()
    parent_no = tasks[0]["No"]

    p1 = store.add_task({"タスク": "子1", "__parent_id": tasks[0]["__kanban_id"], "ステータス": "完了"})
    p2 = store.add_task({"タスク": "子2", "__parent_id": tasks[0]["__kanban_id"], "ステータス": "完了"})

    tasks_after = store.get_tasks()
    parent = next(t for t in tasks_after if t["No"] == parent_no)
    assert parent["ステータス"] == "完了"


def test_auto_sync_parent_status_not_complete_if_child_remaining(tmp_path):
    excel_path = tmp_path / "board.xlsx"
    _build_initial_workbook(excel_path, "Sheet1")
    store = TaskStore(excel_path)

    tasks = store.get_tasks()
    parent_no = tasks[0]["No"]

    store.add_task({"タスク": "子1", "__parent_id": tasks[0]["__kanban_id"], "ステータス": "完了"})
    store.add_task({"タスク": "子2", "__parent_id": tasks[0]["__kanban_id"], "ステータス": "進行中"})

    tasks_after = store.get_tasks()
    parent = next(t for t in tasks_after if t["No"] == parent_no)
    assert parent["ステータス"] != "完了"


def test_auto_sync_cascades_to_grandparent(tmp_path):
    excel_path = tmp_path / "board.xlsx"
    _build_initial_workbook(excel_path, "Sheet1")
    store = TaskStore(excel_path)

    tasks = store.get_tasks()
    grandparent_id = tasks[0]["__kanban_id"]

    parent = store.add_task({"タスク": "親タスク", "__parent_id": grandparent_id, "ステータス": "未着手"})

    store.add_task({"タスク": "孫1", "__parent_id": parent["__kanban_id"], "ステータス": "完了"})
    store.add_task({"タスク": "孫2", "__parent_id": parent["__kanban_id"], "ステータス": "完了"})

    tasks_after = store.get_tasks()
    gp = next(t for t in tasks_after if t["__kanban_id"] == grandparent_id)
    assert gp["ステータス"] == "完了"
