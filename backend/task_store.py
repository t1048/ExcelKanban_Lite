# -*- coding: utf-8 -*-
from __future__ import annotations

import datetime as dt
import os
import shutil
import threading
import uuid
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation

from .constants import (
    DEFAULT_STATUSES,
    DEFAULT_PRIORITY_LEVELS,
    DEFAULT_VALIDATIONS,
    HIDDEN_META_COLUMNS,
    META_ID_COLUMN,
    PARENT_ID_COLUMN,
    TASK_COLUMNS,
)
from .utils import (
    format_priority,
    from_iso_date_str,
    normalize_priority,
    sanitize_text_for_excel,
    to_iso_date_str,
)

MAX_BACKUP_COUNT = 10


class TaskStore:
    def __init__(self, excel_path: Path, sheet_name: str | None = None):
        self.excel_path = excel_path
        self._lock = threading.RLock()
        self._meta_id_column = META_ID_COLUMN
        self._meta_parent_id_column = PARENT_ID_COLUMN
        self._df = pd.DataFrame(columns=TASK_COLUMNS + HIDDEN_META_COLUMNS)
        self._column_order: List[str] = list(TASK_COLUMNS) + list(HIDDEN_META_COLUMNS)
        self._statuses: List[str] = list(DEFAULT_STATUSES)
        self._requested_sheet_name: str | None = (
            sheet_name.strip() if isinstance(sheet_name, str) and sheet_name.strip() else None
        )
        self._sheet_name: str | None = self._requested_sheet_name
        self._validations: Dict[str, List[str]] = {
            key: list(values) for key, values in DEFAULT_VALIDATIONS.items()
        }
        self._last_saved_at: Optional[dt.datetime] = None
        self._last_saved_mtime: Optional[float] = None
        self._last_loaded_mtime: Optional[float] = None
        self._dirty_row_ids: Set[str] = set()
        self._deleted_row_ids: Set[str] = set()
        self._last_saved_snapshot: pd.DataFrame = pd.DataFrame(columns=self._df.columns)
        self.load_excel()

    def load_excel(self):
        with self._lock:
            requested_sheet = self._requested_sheet_name
            if not self.excel_path.exists():
                df = pd.DataFrame(columns=TASK_COLUMNS)
                df.to_excel(
                    self.excel_path,
                    index=False,
                    sheet_name=requested_sheet or "Sheet1",
                )

            wb = load_workbook(self.excel_path, data_only=False)
            try:
                if requested_sheet:
                    if requested_sheet in wb.sheetnames:
                        sheet_name = requested_sheet
                    else:
                        ws = wb.create_sheet(title=requested_sheet)
                        ws.append(TASK_COLUMNS)
                        wb.save(self.excel_path)
                        sheet_name = requested_sheet
                else:
                    sheet_name = wb.sheetnames[0]
                ws = wb[sheet_name]
                self._sheet_name = sheet_name
                extracted = self._extract_validations(wb, ws)
                validations: Dict[str, List[str]] = {
                    key: list(values) for key, values in DEFAULT_VALIDATIONS.items()
                }
                for column, values in extracted.items():
                    if not values:
                        continue
                    validations[column] = list(values)
                self._validations = validations

                df = pd.read_excel(
                    self.excel_path,
                    sheet_name=sheet_name,
                    engine="openpyxl",
                    index_col=None,
                )

                if "No" in df.columns:
                    df = df.drop(columns=["No"])

                for col in TASK_COLUMNS:
                    if col not in df.columns:
                        df[col] = pd.NA

                extra_columns = [col for col in df.columns if col not in TASK_COLUMNS]
                ordered_columns = TASK_COLUMNS + extra_columns
                df = df[ordered_columns].copy()

                # 大分類・中分類 → 親子階層の自動移行
                if "大分類" in df.columns or "中分類" in df.columns:
                    try:
                        self._rotate_backups()
                        backup_path = self.excel_path.parent / (
                            f"{self.excel_path.stem}.bak_migrate_{dt.datetime.now().strftime('%Y%m%d_%H%M%S')}{self.excel_path.suffix}"
                        )
                        shutil.copy2(self.excel_path, backup_path)
                        df = self._migrate_categories_to_hierarchy(df)
                        # 移行結果をExcelに書き戻し
                        wb_migration = load_workbook(self.excel_path)
                        try:
                            ws_mig = wb_migration[self._sheet_name or wb_migration.sheetnames[0]]
                            ws_mig.delete_rows(1, ws_mig.max_row)
                            cols_to_write = list(df.columns)
                            ws_mig.append(cols_to_write)
                            for col_idx, col_name in enumerate(cols_to_write, start=1):
                                if col_name in HIDDEN_META_COLUMNS:
                                    from openpyxl.utils import get_column_letter as _gl
                                    ws_mig.column_dimensions[_gl(col_idx)].hidden = True
                            for _, row_data in df.iterrows():
                                values = []
                                for col_name in cols_to_write:
                                    val = row_data[col_name]
                                    if pd.isna(val):
                                        values.append(None)
                                    else:
                                        values.append(val)
                                ws_mig.append(values)
                            wb_migration.save(self.excel_path)
                        finally:
                            try:
                                wb_migration.close()
                            except Exception:
                                pass
                    except Exception as exc:
                        print(f"[kanban] Category migration failed: {exc}")

                df = df.dropna(how="all", subset=TASK_COLUMNS).reset_index(drop=True)

                df["タスク"] = df["タスク"].apply(
                    lambda v: "" if pd.isna(v) else str(v).strip()
                )
                df = df[df["タスク"] != ""].reset_index(drop=True)

                text_columns = ["ステータス", "担当者", "備考"]
                for col in text_columns:
                    if col in df.columns:
                        df[col] = df[col].apply(lambda v: "" if pd.isna(v) else str(v).strip())

                if "期限" in df.columns:
                    df["期限"] = pd.to_datetime(df["期限"], errors="coerce").dt.date.astype(object)

                had_meta_column = self._meta_id_column in df.columns

                df = self._ensure_row_ids(df)
                persisted_mtime: Optional[float] = None
                if not had_meta_column:
                    persisted_mtime = self._persist_row_ids_to_workbook(df)
                    if persisted_mtime is not None:
                        self._last_saved_at = dt.datetime.now()

                local_df = self._ensure_row_ids(self._df.copy())

                excel_columns = list(df.columns)
                local_columns = list(local_df.columns)
                all_columns = list(
                    dict.fromkeys(excel_columns + [col for col in local_columns if col not in excel_columns])
                )

                dirty_ids = set(self._dirty_row_ids)
                deleted_ids = set(self._deleted_row_ids)

                excel_row_ids: List[str] = []
                excel_row_map: Dict[str, pd.Series] = {}
                for _, excel_row in df.iterrows():
                    row_id = str(excel_row[self._meta_id_column])
                    excel_row_ids.append(row_id)
                    excel_row_map[row_id] = excel_row

                local_row_map: Dict[str, pd.Series] = {}
                for _, local_row in local_df.iterrows():
                    row_id = str(local_row[self._meta_id_column])
                    local_row_map[row_id] = local_row

                merged_rows: List[pd.Series] = []
                new_dirty_ids: Set[str] = set()

                for row_id in excel_row_ids:
                    if row_id in deleted_ids:
                        continue
                    if row_id in local_row_map and row_id in dirty_ids:
                        source_row = local_row_map[row_id]
                        new_dirty_ids.add(row_id)
                    else:
                        source_row = excel_row_map[row_id]
                    merged_rows.append(source_row.reindex(all_columns))

                for _, local_row in local_df.iterrows():
                    row_id = str(local_row[self._meta_id_column])
                    if row_id in deleted_ids or row_id in excel_row_map:
                        continue
                    if row_id not in dirty_ids:
                        continue
                    merged_rows.append(local_row.reindex(all_columns))
                    new_dirty_ids.add(row_id)

                if merged_rows:
                    merged_df = pd.DataFrame(merged_rows)
                    merged_df = merged_df.reindex(columns=all_columns)
                else:
                    merged_df = pd.DataFrame(columns=all_columns)

                new_excel_columns = [
                    col
                    for col in all_columns
                    if col in df.columns and col not in local_df.columns and col != self._meta_id_column
                ]
                for col in new_excel_columns:
                    merged_df[col] = merged_df[col].where(~merged_df[col].isna(), pd.NA)

                merged_df = merged_df.dropna(how="all", subset=TASK_COLUMNS).reset_index(drop=True)

                merged_df["タスク"] = merged_df["タスク"].apply(
                    lambda v: "" if pd.isna(v) else str(v).strip()
                )
                merged_df = merged_df[merged_df["タスク"] != ""].reset_index(drop=True)

                if "期限" in merged_df.columns:
                    merged_df["期限"] = pd.to_datetime(merged_df["期限"], errors="coerce").dt.date.astype(object)

                extra_columns = [col for col in merged_df.columns if col not in TASK_COLUMNS]
                ordered_columns = TASK_COLUMNS + extra_columns
                merged_df = merged_df[ordered_columns]
                self._column_order = ordered_columns

                self._df = merged_df
                self._dirty_row_ids = new_dirty_ids
                excel_id_set = set(excel_row_ids)
                self._deleted_row_ids = {row_id for row_id in deleted_ids if row_id in excel_id_set}

                self._last_saved_snapshot = df.reindex(columns=all_columns).copy(deep=True)
                self._last_saved_snapshot = self._last_saved_snapshot.where(
                    ~self._last_saved_snapshot.isna(), pd.NA
                )

                if self._validations.get("ステータス"):
                    base = list(self._validations["ステータス"])
                    extras = [
                        s
                        for s in self._df["ステータス"].dropna().astype(str).tolist()
                        if s and s not in base
                    ]
                    self._statuses = base + extras
                else:
                    self._rebuild_statuses_from_df(self._df)
                mtime = persisted_mtime if persisted_mtime is not None else self._get_file_mtime()
                self._last_loaded_mtime = mtime
                if persisted_mtime is not None:
                    self._last_saved_mtime = mtime
                elif self._last_saved_mtime is None:
                    self._last_saved_mtime = self._last_loaded_mtime
            finally:
                try:
                    wb.close()
                except Exception:
                    pass

    def _generate_row_id(self) -> str:
        return uuid.uuid4().hex

    def _ensure_meta_columns(self, df: pd.DataFrame) -> pd.DataFrame:
        if self._meta_id_column not in df.columns:
            df[self._meta_id_column] = pd.NA
        if self._meta_parent_id_column not in df.columns:
            df[self._meta_parent_id_column] = pd.NA
        return df

    def _migrate_categories_to_hierarchy(self, df: pd.DataFrame) -> pd.DataFrame:
        """大分類・中分類列を親子階層(__parent_id)に移行する。"""
        has_major = "大分類" in df.columns
        has_minor = "中分類" in df.columns
        has_parent = self._meta_parent_id_column in df.columns

        if has_parent or (not has_major and not has_minor):
            cols_to_drop = [c for c in ["大分類", "中分類"] if c in df.columns]
            if cols_to_drop:
                df = df.drop(columns=cols_to_drop)
            return df

        # バックアップは呼び出し側(load_excel)で実施済み
        major_values = df["大分類"].apply(
            lambda v: "" if pd.isna(v) else str(v).strip()
        )
        unique_majors = sorted(set(v for v in major_values if v))

        # 大分類ごとに親タスク行を作成
        parent_rows: List[Dict[str, Any]] = []
        parent_id_map: Dict[str, str] = {}
        for major_name in unique_majors:
            parent_id = self._generate_row_id()
            parent_id_map[major_name] = parent_id
            parent_rows.append({
                "ステータス": "未着手",
                "タスク": major_name,
                "担当者": "",
                "優先度": "",
                "期限": pd.NaT,
                "備考": "",
                self._meta_id_column: parent_id,
                self._meta_parent_id_column: pd.NA,
            })

        # 元のタスクに__parent_idを設定し、中分類を備考に退避
        parent_ids: List[Any] = []
        # 備考列の型を文字列に統一（NaN混在時のfloat64問題回避）
        if "備考" in df.columns:
            df["備考"] = df["備考"].apply(lambda v: "" if pd.isna(v) else str(v))
        for _, row in df.iterrows():
            major = str(row.get("大分類", "") or "").strip()
            minor = str(row.get("中分類", "") or "").strip()
            pid = parent_id_map.get(major, pd.NA)
            parent_ids.append(pid if pid else pd.NA)

            # 中分類を備考の先頭に退避
            if minor:
                existing_note = str(row.get("備考", "") or "")
                if existing_note == "nan":
                    existing_note = ""
                prefix = f"[旧中分類: {minor}]"
                if prefix not in existing_note:
                    new_note = f"{prefix} {existing_note}".strip()
                    df.at[row.name, "備考"] = new_note

        df[self._meta_parent_id_column] = parent_ids

        # 大分類・中分類列を削除
        cols_to_drop = [c for c in ["大分類", "中分類"] if c in df.columns]
        if cols_to_drop:
            df = df.drop(columns=cols_to_drop)

        # 親タスク行を追加
        if parent_rows:
            parent_df = pd.DataFrame(parent_rows)
            # 列を揃える
            for col in df.columns:
                if col not in parent_df.columns:
                    parent_df[col] = pd.NA
            parent_df = parent_df[df.columns]
            df = pd.concat([parent_df, df], ignore_index=True)

        return df

    def _check_circular_ref(self, task_id: str, new_parent_id: str) -> bool:
        """new_parent_idがtask_idの子孫であればTrue（循環参照）。"""
        if not new_parent_id or not task_id:
            return False
        if task_id == new_parent_id:
            return True
        visited = set()
        current = new_parent_id
        while current:
            if current == task_id:
                return True
            if current in visited:
                break
            visited.add(current)
            parent_row = self._df[self._df[self._meta_id_column] == current]
            if parent_row.empty:
                break
            parent_val = parent_row.iloc[0].get(self._meta_parent_id_column, "")
            if parent_val is None or pd.isna(parent_val) or str(parent_val).strip() == "":
                break
            current = str(parent_val).strip()
        return False

    def get_children(self, task_id: str) -> List[Dict[str, Any]]:
        """指定タスクの直接の子タスク一覧を返す。"""
        with self._lock:
            mask = self._df[self._meta_parent_id_column] == task_id
            children = self._df[mask]
            return [
                self._format_row(i, children.iloc[j])
                for j, i in enumerate(children.index)
            ]

    def get_descendant_ids(self, task_id: str) -> List[str]:
        """指定タスクの全子孫IDを再帰的に収集。"""
        result: List[str] = []
        queue = [task_id]
        while queue:
            current = queue.pop(0)
            mask = self._df[self._meta_parent_id_column] == current
            child_ids = self._df.loc[mask, self._meta_id_column].tolist()
            for cid in child_ids:
                cid_str = str(cid).strip()
                if cid_str and cid_str not in result:
                    result.append(cid_str)
                    queue.append(cid_str)
        return result

    def get_ancestor_ids(self, task_id: str) -> List[str]:
        """指定タスクの全祖先IDを返す。"""
        result: List[str] = []
        current = task_id
        visited = set()
        while current:
            if current in visited:
                break
            visited.add(current)
            row = self._df[self._df[self._meta_id_column] == current]
            if row.empty:
                break
            parent_val = row.iloc[0].get(self._meta_parent_id_column, "")
            if parent_val is None or pd.isna(parent_val) or str(parent_val).strip() == "":
                break
            parent_id = str(parent_val).strip()
            result.append(parent_id)
            current = parent_id
        return result

    def get_task_tree(self) -> List[Dict[str, Any]]:
        """タスクをツリー構造(children配列付き)で返す。"""
        with self._lock:
            tasks = [self._format_row(i, self._df.iloc[i]) for i in range(len(self._df))]
            id_map: Dict[str, Dict[str, Any]] = {}
            for t in tasks:
                t["children"] = []
                id_map[t["__kanban_id"]] = t
            roots: List[Dict[str, Any]] = []
            for t in tasks:
                pid = t.get("__parent_id", "")
                if pid and pid in id_map:
                    id_map[pid]["children"].append(t)
                else:
                    roots.append(t)
            return roots

    def _ensure_row_ids(self, df: pd.DataFrame) -> pd.DataFrame:
        df = self._ensure_meta_columns(df.copy())
        for idx, value in df[self._meta_id_column].items():
            if value is None or (isinstance(value, str) and not value.strip()) or pd.isna(value):
                df.at[idx, self._meta_id_column] = self._generate_row_id()
            else:
                df.at[idx, self._meta_id_column] = str(value).strip()
        return df

    def _persist_row_ids_to_workbook(self, df: pd.DataFrame) -> Optional[float]:
        sheet_name = self._sheet_name
        if not sheet_name:
            return None
        try:
            wb = load_workbook(self.excel_path)
        except FileNotFoundError:
            return None
        try:
            if sheet_name not in wb.sheetnames:
                return None
            ws = wb[sheet_name]

            header_row = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ()))
            id_col_idx: Optional[int] = None
            for idx, value in enumerate(header_row, start=1):
                if isinstance(value, str) and value == self._meta_id_column:
                    id_col_idx = idx
                    break

            changes_made = False
            if id_col_idx is None:
                id_col_idx = len(header_row) + 1 if header_row else 1
                ws.cell(row=1, column=id_col_idx, value=self._meta_id_column)
                changes_made = True

            header_map: Dict[str, int] = {}
            for col_idx in range(1, ws.max_column + 1):
                header_value = ws.cell(row=1, column=col_idx).value
                if isinstance(header_value, str):
                    header_map[header_value] = col_idx

            column_letter = get_column_letter(id_col_idx)
            if ws.column_dimensions[column_letter].hidden is not True:
                ws.column_dimensions[column_letter].hidden = True
                changes_made = True

            id_values = [str(value) for value in df[self._meta_id_column].tolist()]
            task_col_idx = header_map.get("タスク")
            id_index = 0
            for row_idx in range(2, ws.max_row + 1):
                if id_index >= len(id_values):
                    break
                if task_col_idx is not None:
                    task_cell = ws.cell(row=row_idx, column=task_col_idx).value
                    task_text = "" if task_cell is None else str(task_cell).strip()
                    if not task_text:
                        continue
                ws.cell(row=row_idx, column=id_col_idx).value = id_values[id_index]
                id_index += 1
                changes_made = True

            if not changes_made:
                return None

            wb.save(self.excel_path)
            try:
                return self._get_file_mtime()
            finally:
                pass
        finally:
            try:
                wb.close()
            except Exception:
                pass

    def _get_row_id_at_index(self, row_index: int) -> str:
        self._df = self._ensure_meta_columns(self._df)
        value = self._df.at[row_index, self._meta_id_column]
        if value is None or (isinstance(value, str) and not value.strip()) or pd.isna(value):
            new_id = self._generate_row_id()
            self._df.at[row_index, self._meta_id_column] = new_id
            return new_id
        return str(value).strip()

    def _get_file_mtime(self) -> Optional[float]:
        try:
            return self.excel_path.stat().st_mtime
        except FileNotFoundError:
            return None

    def get_last_saved_markers(self) -> Tuple[Optional[dt.datetime], Optional[float]]:
        with self._lock:
            return self._last_saved_at, self._last_saved_mtime

    def _rebuild_statuses_from_df(self, df: pd.DataFrame | None = None):
        if df is None:
            df = self._df
        status_values = [
            s for s in df["ステータス"].dropna().astype(str).tolist() if s
        ]
        merged: List[str] = []
        for name in status_values + DEFAULT_STATUSES:
            if name not in merged:
                merged.append(name)
        self._statuses = merged

    def _extract_validations(self, wb, ws) -> Dict[str, List[str]]:
        validations: Dict[str, List[str]] = {}
        dv_list = getattr(ws, "data_validations", None)
        if not dv_list:
            return validations

        for dv in getattr(dv_list, "dataValidation", []) or []:
            if dv.type != "list":
                continue
            values = self._resolve_validation_values(wb, ws, dv)
            if not values:
                continue
            try:
                ranges = list(dv.ranges)
            except TypeError:
                ranges = []
            for cell_range in ranges:
                min_col, min_row, max_col, max_row = range_boundaries(str(cell_range))
                for col_idx in range(min_col, max_col + 1):
                    header = ws.cell(row=1, column=col_idx).value
                    if isinstance(header, str) and header in TASK_COLUMNS:
                        validations[header] = list(values)
        return validations

    def _resolve_validation_values(self, wb, ws, dv) -> List[str]:
        formula = (dv.formula1 or "").strip()
        if not formula:
            return []
        if formula.startswith('"') and formula.endswith('"'):
            content = formula[1:-1]
            parts = [p.replace('""', '"').strip() for p in content.split(",")]
            return [p for p in parts if p]
        if formula.startswith("="):
            target = formula[1:]
            sheet_name = ws.title
            if "!" in target:
                sheet_part, range_part = target.split("!", 1)
                sheet_name = sheet_part.strip()
                if sheet_name.startswith("'") and sheet_name.endswith("'"):
                    sheet_name = sheet_name[1:-1].replace("''", "'")
            else:
                range_part = target
            range_part = range_part.strip()
            try:
                target_ws = wb[sheet_name]
            except KeyError:
                return []
            values: List[str] = []
            for row in target_ws[range_part]:
                cells = row if isinstance(row, (list, tuple)) else (row,)
                for cell in cells:
                    if cell.value is None:
                        continue
                    text = str(cell.value).strip()
                    if not text:
                        continue
                    if text not in values:
                        values.append(text)
            return values
        return []

    def _to_excel_value(self, col_name: str, value: Any):
        if value is None or value is pd.NA:
            return None
        try:
            if pd.isna(value):
                return None
        except Exception:
            pass

        if isinstance(value, pd.Timestamp):
            return value.to_pydatetime()
        if isinstance(value, dt.datetime):
            return value
        if isinstance(value, dt.date):
            return value

        if col_name == "期限":
            if isinstance(value, str):
                parsed = from_iso_date_str(value)
                return None if parsed is pd.NaT else parsed

        if isinstance(value, str):
            text = sanitize_text_for_excel(value)
            if text is None:
                return None
            return text

        return value

    def _build_validation_formula(self, values: List[str]) -> str:
        escaped = [v.replace('"', '""') for v in values]
        return '"' + ",".join(escaped) + '"'

    def get_tasks(self) -> List[Dict[str, Any]]:
        with self._lock:
            return [self._format_row(i, self._df.iloc[i]) for i in range(len(self._df))]

    def get_statuses(self) -> List[str]:
        with self._lock:
            return list(self._statuses)

    def get_validations(self) -> Dict[str, List[str]]:
        with self._lock:
            return {k: list(v) for k, v in self._validations.items()}

    def get_state_snapshot(self) -> Dict[str, Any]:
        with self._lock:
            tasks = [self._format_row(i, self._df.iloc[i]) for i in range(len(self._df))]
            statuses = list(self._statuses)
            validations = {k: list(v) for k, v in self._validations.items()}
            return {
                "tasks": tasks,
                "statuses": statuses,
                "validations": validations,
            }

    def set_validations(self, mapping: Dict[str, List[Any]]):
        with self._lock:
            cleaned: Dict[str, List[str]] = {}
            for col in TASK_COLUMNS:
                raw_values = mapping.get(col)
                if not raw_values:
                    continue
                values: List[str] = []
                for value in raw_values:
                    text = str(value or "").strip()
                    if text and text not in values:
                        values.append(text)
                if values:
                    cleaned[col] = values
            merged: Dict[str, List[str]] = {
                key: list(values) for key, values in DEFAULT_VALIDATIONS.items()
            }
            merged.update(cleaned)
            self._validations = merged
            if self._validations.get("ステータス"):
                base = list(self._validations["ステータス"])
                extras = [
                    s
                    for s in self._df["ステータス"].dropna().astype(str).tolist()
                    if s and s not in base
                ]
                self._statuses = base + extras
            else:
                self._rebuild_statuses_from_df()

    def _format_row(self, idx: int, row: pd.Series) -> Dict[str, Any]:
        parent_id = row.get(self._meta_parent_id_column)
        return {
            "No": int(idx + 1),
            "ステータス": "" if pd.isna(row["ステータス"]) else str(row["ステータス"]),
            "タスク": "" if pd.isna(row["タスク"]) else str(row["タスク"]),
            "担当者": "" if pd.isna(row["担当者"]) else str(row["担当者"]),
            "優先度": format_priority(row["優先度"]),
            "期限": to_iso_date_str(row["期限"]),
            "備考": "" if pd.isna(row["備考"]) else str(row["備考"]),
            "__kanban_id": str(row.get(self._meta_id_column, "") or "").strip(),
            "__parent_id": "" if (parent_id is None or pd.isna(parent_id) or str(parent_id).strip() == "") else str(parent_id).strip(),
        }

    def _resolve_row_index(self, no_value: int) -> int:
        try:
            no = int(no_value)
        except (TypeError, ValueError):
            raise KeyError(f"No={no_value} は存在しません。")
        if no <= 0 or no > len(self._df):
            raise KeyError(f"No={no_value} は存在しません。")
        return no - 1

    def _ensure_status_registered(self, name: str):
        name = name.strip()
        if name and name not in self._statuses:
            self._statuses.append(name)

    def _auto_sync_parent_status(self, child_id: str):
        row = self._df[self._df[self._meta_id_column] == child_id]
        if row.empty:
            return
        parent_val = row.iloc[0].get(self._meta_parent_id_column)
        parent_id = "" if parent_val is None or pd.isna(parent_val) else str(parent_val).strip()
        if not parent_id:
            return
        siblings = self._df[self._df[self._meta_parent_id_column] == parent_id]
        if len(siblings) == 0:
            return
        all_completed = all(
            str(s.get("ステータス", "") or "").strip() == "完了"
            for _, s in siblings.iterrows()
        )
        parent_row = self._df[self._df[self._meta_id_column] == parent_id]
        if parent_row.empty:
            return
        current_parent_status = str(parent_row.iloc[0].get("ステータス", "") or "").strip()

        if all_completed and current_parent_status != "完了":
            self._df.loc[self._df[self._meta_id_column] == parent_id, "ステータス"] = "完了"
            self._ensure_status_registered("完了")
        elif not all_completed and current_parent_status == "完了":
            self._df.loc[self._df[self._meta_id_column] == parent_id, "ステータス"] = "進行中"
            self._ensure_status_registered("進行中")
        else:
            return

        parent_row_id = str(parent_row.iloc[0][self._meta_id_column])
        self._dirty_row_ids.add(parent_row_id)
        self._auto_sync_parent_status(parent_id)

    def add_task(self, payload: Dict[str, Any]) -> Dict[str, Any]:
        with self._lock:
            task = dict(payload)

            status = str(task.get("ステータス", "") or "").strip()
            title = str(task.get("タスク", "") or "").strip()
            assignee = str(task.get("担当者", "") or "").strip()
            notes = str(task.get("備考", "") or "")
            due = from_iso_date_str(task.get("期限", ""))
            priority = normalize_priority(task.get("優先度"))
            parent_id = str(task.get("__parent_id", "") or task.get("parent_id", "") or "").strip()

            if not title:
                raise ValueError("タスクは必須項目です。")

            if parent_id:
                # 循環参照チェック（新規タスクなので親のみチェック）
                parent_exists = not self._df[self._df[self._meta_id_column] == parent_id].empty
                if not parent_exists:
                    raise ValueError(f"親タスクID {parent_id} が見つかりません。")

            row = {
                "ステータス": status,
                "タスク": title,
                "担当者": assignee,
                "優先度": priority,
                "期限": due,
                "備考": notes,
            }

            new_index = len(self._df)
            self._df = self._ensure_meta_columns(self._df)
            full_row = {col: pd.NA for col in self._df.columns}
            full_row.update(row)
            row_id = self._generate_row_id()
            full_row[self._meta_id_column] = row_id
            if parent_id:
                full_row[self._meta_parent_id_column] = parent_id
            self._df.loc[new_index] = full_row
            self._ensure_status_registered(status)
            self._dirty_row_ids.add(row_id)
            self._deleted_row_ids.discard(row_id)
            if parent_id:
                self._auto_sync_parent_status(row_id)
            return self._format_row(new_index, self._df.iloc[new_index])

    def update_task(self, no_value: int, patch: Dict[str, Any]) -> Dict[str, Any]:
        with self._lock:
            i = self._resolve_row_index(int(no_value))
            row_index = self._df.index[i]

            updates: Dict[str, Any] = {}

            if "ステータス" in patch:
                updates["ステータス"] = str(patch["ステータス"] or "").strip()
            if "タスク" in patch:
                title = str(patch["タスク"] or "").strip()
                if not title:
                    raise ValueError("タスクは必須項目です。")
                updates["タスク"] = title
            if "担当者" in patch:
                updates["担当者"] = str(patch["担当者"] or "").strip()
            if "優先度" in patch:
                updates["優先度"] = normalize_priority(patch["優先度"])
            if "期限" in patch:
                updates["期限"] = from_iso_date_str(patch["期限"])
            if "備考" in patch:
                updates["備考"] = str(patch["備考"] or "")

            # __parent_id の更新
            if "__parent_id" in patch or "parent_id" in patch:
                new_parent_id = str(patch.get("__parent_id", patch.get("parent_id", "")) or "").strip()
                current_id = self._get_row_id_at_index(row_index)
                if new_parent_id:
                    # 循環参照チェック
                    if self._check_circular_ref(current_id, new_parent_id):
                        raise ValueError("循環参照が検出されました。親子関係を変更できません。")
                    # 親タスクの存在チェック
                    parent_exists = not self._df[self._df[self._meta_id_column] == new_parent_id].empty
                    if not parent_exists:
                        raise ValueError(f"親タスクID {new_parent_id} が見つかりません。")
                updates[self._meta_parent_id_column] = new_parent_id if new_parent_id else pd.NA

            if "ステータス" in updates:
                self._ensure_status_registered(updates["ステータス"])

            for column, value in updates.items():
                self._df.at[row_index, column] = value

            row_id = self._get_row_id_at_index(row_index)
            self._dirty_row_ids.add(row_id)
            self._deleted_row_ids.discard(row_id)

            if "ステータス" in updates:
                self._auto_sync_parent_status(row_id)

            return self._format_row(i, self._df.loc[row_index])

    def move_task(self, no_value: int, new_status: str) -> Dict[str, Any]:
        return self.update_task(int(no_value), {"ステータス": new_status})

    def delete_task(self, no_value: int, cascade: bool = True) -> bool:
        with self._lock:
            try:
                idx = self._resolve_row_index(int(no_value))
            except KeyError:
                return False
            row_index = self._df.index[idx]
            row_id = self._get_row_id_at_index(row_index)

            if cascade:
                descendant_ids = self.get_descendant_ids(row_id)
                if descendant_ids:
                    self._df = self._df[
                        ~self._df[self._meta_id_column].isin(descendant_ids)
                    ].reset_index(drop=True)
            else:
                # 子タスクの__parent_idをクリアしてルートに昇格
                self._df.loc[
                    self._df[self._meta_parent_id_column] == row_id,
                    self._meta_parent_id_column,
                ] = pd.NA

            self._df = self._df[self._df[self._meta_id_column] != row_id].reset_index(drop=True)
            self._dirty_row_ids.discard(row_id)
            self._deleted_row_ids.add(row_id)
            return True

    def _rotate_backups(self) -> None:
        pattern = f"{self.excel_path.stem}.bak_*{self.excel_path.suffix}"
        backups = sorted(
            [p for p in self.excel_path.parent.glob(pattern)],
            key=lambda p: p.stat().st_mtime,
        )
        while len(backups) >= MAX_BACKUP_COUNT:
            try:
                backups[0].unlink()
                backups.pop(0)
            except OSError:
                break

    def save_excel(self) -> str:
        with self._lock:
            ts = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
            backup_path = self.excel_path.parent / (
                f"{self.excel_path.stem}.bak_{ts}{self.excel_path.suffix}"
            )
            tmp_path = self.excel_path.with_name(
                f"{self.excel_path.stem}.tmp_{ts}{self.excel_path.suffix}"
            )

            df = self._df.copy()
            df = df.drop(columns=[c for c in ["大分類", "中分類"] if c in df.columns])
            wb = None
            try:
                if self.excel_path.exists():
                    wb = load_workbook(self.excel_path)
                else:
                    wb = Workbook()
                if not self._sheet_name:
                    self._sheet_name = wb.sheetnames[0]
                if self._sheet_name not in wb.sheetnames:
                    ws = wb.create_sheet(title=self._sheet_name)
                else:
                    ws = wb[self._sheet_name]

                current_columns = list(df.columns)
                extra_columns = [col for col in current_columns if col not in TASK_COLUMNS]
                ordered_columns = TASK_COLUMNS + extra_columns
                df = df[ordered_columns]
                self._column_order = ordered_columns

                df_with_no = df.copy()
                df_with_no.insert(0, 'No', range(1, len(df_with_no) + 1))
                excel_columns = ['No'] + ordered_columns

                ws.delete_rows(1, ws.max_row)
                ws.append(excel_columns)
                for idx, col_name in enumerate(excel_columns, start=1):
                    if col_name in HIDDEN_META_COLUMNS:
                        col_letter = get_column_letter(idx)
                        ws.column_dimensions[col_letter].hidden = True
                for row_num, row in enumerate(df_with_no.itertuples(index=False, name=None), start=1):
                    values: List[Any] = []
                    for col_name, value in zip(excel_columns, row):
                        if col_name == 'No':
                            values.append(row_num)
                        else:
                            values.append(self._to_excel_value(col_name, value))
                    ws.append(values)

                try:
                    due_col_idx = TASK_COLUMNS.index("期限") + 2
                except ValueError:
                    due_col_idx = None

                if due_col_idx is not None and ws.max_row >= 2:
                    for cell in ws.iter_rows(
                        min_row=2,
                        max_row=ws.max_row,
                        min_col=due_col_idx,
                        max_col=due_col_idx,
                    ):
                        for target in cell:
                            target.number_format = "yyyy/mm/dd"

                if ws.data_validations is not None:
                    ws.data_validations.dataValidation = []

                for col_name, values in self._validations.items():
                    if col_name not in TASK_COLUMNS or not values:
                        continue
                    try:
                        idx = TASK_COLUMNS.index(col_name) + 2
                    except ValueError:
                        continue
                    col_letter = get_column_letter(idx)
                    formula = self._build_validation_formula(values)
                    dv = DataValidation(
                        type="list",
                        formula1=formula,
                        allow_blank=True,
                        showDropDown=False,
                    )
                    dv.add(f"${col_letter}$2:${col_letter}$1048576")
                    ws.add_data_validation(dv)

                wb.save(tmp_path)
                if tmp_path.exists():
                    os.replace(tmp_path, self.excel_path)
                shutil.copy2(self.excel_path, backup_path)
                self._rotate_backups()
                self._last_saved_at = dt.datetime.now()
                self._last_saved_mtime = self._get_file_mtime()
                self._last_loaded_mtime = self._last_saved_mtime
                self._dirty_row_ids.clear()
                self._deleted_row_ids.clear()
                self._last_saved_snapshot = df.copy(deep=True)
                self._df = df.copy()
            except Exception:
                if tmp_path.exists():
                    try:
                        tmp_path.unlink()
                    except Exception:
                        pass
                raise
            finally:
                if wb is not None:
                    try:
                        wb.close()
                    except Exception:
                        pass
            return str(self.excel_path.resolve())
